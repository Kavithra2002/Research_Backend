"""
Generate COMB financial model Excel (FS sheet) from MongoDB financial_tables.

Fills annual figures 2017-2025 for Commercial Bank of Ceylon PLC using the
client's COMB model.xlsx layout. Writes a JSON report of missing line items.
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.styles import PatternFill
from pymongo import MongoClient

from comb_cell_status import STATUS_CONFIRMED_ABSENT, STATUS_TEMPLATE_MISMATCH
from comb_workbook_store import (
    CombWorkbookStore,
    baseline_export_years,
    build_note_metadata,
    resolve_export_years,
)
from comb_template_reference import (
    apply_quarterly_template_validation,
    apply_template_validation_to_rows,
)
from extraction_aliases_store import patterns_for_label

SCRIPT_DIR = Path(__file__).resolve().parent
BACKEND_DIR = SCRIPT_DIR.parent
TEMPLATE = BACKEND_DIR / "New_Updates" / "COMB model - updated.xlsx"
if not TEMPLATE.exists():
    TEMPLATE = BACKEND_DIR / "New_Updates" / "COMB model.xlsx"
OUTPUT = BACKEND_DIR / "New_Updates" / "COMB_model_2017_2025_generated.xlsx"
MISSING_REPORT = BACKEND_DIR / "New_Updates" / "COMB_model_missing_values.json"

COMPANY_SLUG = "Commercial_Bank_of_Ceylon_PLC"
YEARS = baseline_export_years()
EXPORT_YEARS = YEARS
LABEL_COL = 2  # column B holds line-item labels on FS sheet

SKIP_LABELS = {
    "INCOME STATEMENT",
    "OCI",
    "IS check",
    "BALANCE SHEET",
    "Assets",
    "Liabilities",
    "Equity",
    "Earnings per share",
    "Profit attributable to:",
    "Less: Expenses",
    "Memorandum information",
    "Cash flows from investing activities",
    "Cash flows from financing activities",
    "Adjustments for:",
}

SECTION_HEADERS = {
    "INCOME STATEMENT",
    "OCI",
    "BALANCE SHEET",
    "CASH FLOW STATEMENT",
}

SUBSECTION_HEADERS = {
    "Assets",
    "Liabilities",
    "Equity",
    "Memorandum information",
    "Cash flows from investing activities",
    "Cash flows from financing activities",
    "Adjustments for:",
    "Less: Expenses",
}

CHECK_LABELS = {
    "BS check",
}

COMMERCIAL_BANK_SLUG = "Commercial_Bank_of_Ceylon_PLC"
PILOT_YEARS = EXPORT_YEARS  # all template years (2017-2025) may hold extracted data
EMPTY_CELL = "-"
RED_FILL = PatternFill(start_color="FFFFC7CE", end_color="FFFFC7CE", fill_type="solid")
QUARTERLY_EXPORT_YEARS = EXPORT_YEARS
QUARTERLY_PILOT_YEAR = QUARTERLY_EXPORT_YEARS[-1]
QUARTERLY_PILOT_QUARTERS = ["Q1", "Q2", "Q3", "Q4"]
QUARTER_LABELS = {
    "Q1": "Mar",
    "Q2": "Jun",
    "Q3": "Sep",
    "Q4": "Dec",
}


def quarterly_column_defs(years: list[int] | None = None) -> list[dict[str, Any]]:
    """Quarter columns in chronological order (e.g. Mar 2017 … Dec 2024)."""
    year_list = years or QUARTERLY_EXPORT_YEARS
    columns: list[dict[str, Any]] = []
    for year in year_list:
        for quarter in QUARTERLY_PILOT_QUARTERS:
            columns.append(
                {
                    "key": f"{quarter.lower()}_{year}",
                    "label": f"{QUARTER_LABELS[quarter]} {year}",
                    "year": year,
                    "quarter": quarter,
                }
            )
    return columns
EXPORT_SHEETS = ("Cover", "FS", "Drivers", "Ratios", "Quarterly")


def display_value_for_year(year: int, value: float | None) -> float | None:
    """Return stored workbook values for any year column."""
    return value


def _row_year_maps(
    store: CombWorkbookStore,
    sheet: str,
    label: str,
    year_list: list[int],
    *,
    report_type: str = "annual",
    quarter: str | None = None,
    drivers_row: int | None = None,
) -> tuple[dict[str, float | None], dict[str, str]]:
    values: dict[str, float | None] = {}
    statuses: dict[str, str] = {}
    for year in year_list:
        key = str(year)
        raw = store.lookup(
            year,
            label,
            sheet=sheet,
            report_type=report_type,
            quarter=quarter,
            drivers_row=drivers_row,
        )
        values[key] = display_value_for_year(year, raw)
        statuses[key] = store.lookup_status(
            year,
            label,
            sheet=sheet,
            report_type=report_type,
            quarter=quarter,
            drivers_row=drivers_row,
        )
    return values, statuses

# Template label -> ordered DB label search terms (normalized matching)
LABEL_ALIASES: dict[str, list[str]] = {
    "Gross income": ["gross income"],
    "Interest income": ["interest income"],
    "Less: Interest expense": ["less interest expense", "less : interest expense"],
    "Net interest income": ["net interest income"],
    "Fee and commission income": ["fee and commission income"],
    "Less: Fee and commission expense": ["less fee and commission expense"],
    "Net fee and commission income": ["net fee and commission income"],
    "Net gains/(losses) from trading": ["net gains/(losses) from trading"],
    "Net gains/(losses) from derecognition of financial assets": [
        "net gains/(losses) from derecognition of financial assets",
        "net gains/(losses) from financial investments",
    ],
    "Net other operating income": ["net other operating income", "other income (net)"],
    "Total operating income": ["total operating income"],
    "Less: Impairment charges and other losses": [
        "less impairment charges and other losses",
        "less impairment charges for loans and other losses",
        "impairment charges and other losses",
    ],
    "Net operating income": ["net operating income"],
    "Personnel expenses": ["personnel expenses"],
    "Depreciation and amortisation": ["depreciation and amortisation", "depreciation and amortization"],
    "Other operating expenses": ["other operating expenses"],
    "Total operating expenses": ["total operating expenses"],
    "Operating profit before taxes on financial services": [
        "operating profit before taxes on financial services",
        "operating profit before value added tax",
        "operating profit before vat",
    ],
    "Less: Taxes on financial services": [
        "less taxes on financial services",
        "less value added tax",
        "less: taxes on financial services",
    ],
    "Operating profit after taxes on financial services": [
        "operating profit after taxes on financial services",
        "operating profit after value added tax",
        "operating profit after vat",
    ],
    "Share of profit/(loss) of associate, net of tax": [
        "share of profit/(loss) of associate",
        "share of profits of associates",
        "share of profit of associate",
    ],
    "Profit before tax": ["profit before tax"],
    "Less: Income tax expense/(reversal)": ["less income tax expense", "income tax expense"],
    "Profit for the year": ["profit for the year", "profit for the period"],
    "Equity holders of the Bank": ["equity holders of the bank"],
    "Non-controlling interest": ["non-controlling interest", "non controlling interest"],
    "Basic earnings per ordinary share (Rs.)": ["basic earnings per ordinary share", "basic earnings per share"],
    "Diluted earnings per ordinary share (Rs.)": ["diluted earnings per ordinary share", "diluted earnings per share"],
    "Diviend per share": [
        "dividend per share",
        "diviend per share",
        "dividend per share (dps)",
        "dividend per share (dps)(rs.)",
        "dividend per share (dps) (rs.)",
        "dividends - shares",
        "dividends – shares",
    ],
    "Net assets value per ordinary share (Rs.)": [
        "net assets value per ordinary share",
        "net assets value per share",
        "net asset value per share",
        "net assets value per share (rs.)",
    ],
    "Transfer of FV losses reclassification of debt from FVTOCI to AC": [
        "transfer of fair value losses o/a reclassification of debt instruments from fair value through other comprehensive income to amortised cost, net of tax",
        "transfer of fair value losses on reclassification of debt instruments from fair value through other comprehensive income to amortised cost",
        "reclassification of debt instruments from fair value through other comprehensive income to amortised cost",
        "transfer of fv losses reclassification of debt from fvtoci to ac",
    ],
    "Cash and cash equivalents": ["cash and cash equivalents"],
    "Balances with Central Banks": ["balances with central bank", "balances with central banks"],
    "Placements with banks": ["placements with banks", "placements with bank"],
    "Securities purchased under resale agreements": ["securities purchased under resale agreements"],
    "Derivative financial assets": ["derivative financial assets"],
    "Financial assets recognised through profit or loss – measured at fair value": [
        "financial assets recognised through profit or loss",
        "measured at fair value",
        "held for trading",
    ],
    "Financial assets at amortised cost – loans and advances to banks": [
        "loans and advances to banks",
        "loans and receivables to banks",
    ],
    "Financial assets at amortised cost – loans and advances to other customers": [
        "loans and advances to other customers",
        "loans and receivables to other customers",
    ],
    "Financial assets at amortised cost – debt and other financial instruments": [
        "debt and other financial instruments",
        "financial assets at amortised cost - debt",
    ],
    "Financial assets measured at fair value through other comprehensive income": [
        "fair value through other comprehensive income",
        "available for sale",
    ],
    "Investments in subsidiaries": ["investments in subsidiaries"],
    "Investment in associate": ["investment in associate", "investments in associate"],
    "Property, plant and equipment and right-of-use assets": [
        "property, plant and equipment and right-of-use",
        "property, plant & equipment and right-of-use",
    ],
    "Investment properties": ["investment properties", "investment property"],
    "Intangible assets": ["intangible assets"],
    "Deferred tax assets": ["deferred tax assets"],
    "Other assets": ["other assets"],
    "Total assets": ["total assets"],
    "Due to banks": ["due to banks"],
    "Due to other customers": ["due to other customers", "due to customers"],
    "Securities sold under repurchase agreements": ["securities sold under repurchase agreements"],
    "Derivative financial liabilities": ["derivative financial liabilities"],
    "Debt securities issued": ["debt securities issued"],
    "Other liabilities": ["other liabilities"],
    "Total liabilities": ["total liabilities"],
    "Stated capital": ["stated capital"],
    "Reserves": ["reserves"],
    "Retained earnings": ["retained earnings"],
    "Total equity attributable to equity holders of the bank": [
        "total equity attributable to equity holders of the bank",
        "total equity attributable to equity holders",
    ],
    "Non-controlling interests": ["non-controlling interests", "non controlling interests"],
    "Total equity": ["total equity"],
    "Financial liabilities at amortised cost – due to depositors": [
        "due to depositors",
        "financial liabilities at amortised cost due to depositors",
        "deposits from customers",
        "due to other customers deposits from customers",
    ],
    "Financial liabilities at amortised cost – other borrowings": [
        "financial liabilities at amortised cost other borrowings",
        "other borrowings",
    ],
    "Gross cash and cash equivalents as at December 31,": [
        "gross cash and cash equivalents as at december 31",
        "gross cash and cash equivalents",
    ],
    "Less: Impairment charges on cash and cash equivalents": [
        "less impairment charges on cash and cash equivalents",
        "impairment charges on cash and cash equivalents",
    ],
    "Cash and cash equivalents as per Statement of Financial Position": [
        "cash and cash equivalents as per statement of financial position",
    ],
    "Cash and cash equivalents as at January 01,": [
        "cash and cash equivalents as at january 01",
        "cash and cash equivalents as at 1 january",
    ],
    "Dividend paid to shareholders": [
        "dividend paid to shareholders",
        "dividends paid to shareholders",
    ],
    "Dividend paid to non-controlling interest": [
        "dividend paid to non-controlling interest",
        "dividends paid to non-controlling interest",
    ],
    "Net increase/(decrease) in cash and cash equivalents": [
        "net increase/(decrease) in cash and cash equivalents",
        "net increase in cash and cash equivalents",
    ],
    "Payment of lease liabilities/advance payment of right-of-use assets": [
        "payment of lease liabilities",
        "advance payment of right of use assets",
    ],
}

DRIVERS_LABEL_ALIASES: dict[str, list[str]] = {
    "Gross loans and advances": [
        "financial assets at amortised cost loans and advances to other customers",
        "loans and advances to other customers",
        "gross loans and advances",
    ],
    "Less: Provision for impairment": [
        "less provision for impairment",
        "provision for impairment",
        "impairment provision",
        "total impairment provision",
    ],
    "Interest income": ["interest income"],
    "Cash and cash equivalents": ["cash and cash equivalents"],
    "Due to banks": ["due to banks"],
    "Financial liabilities at amortised cost – due to depositors": [
        "due to depositors",
        "financial liabilities at amortised cost due to depositors",
        "deposits from customers",
        "due to other customers deposits from customers",
    ],
}


def load_env() -> dict[str, str]:
    env: dict[str, str] = {}
    path = BACKEND_DIR / ".env"
    if path.exists():
        for raw in path.read_text(encoding="utf-8").splitlines():
            line = raw.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            k, v = line.split("=", 1)
            env[k.strip()] = v.strip().strip("'\"")
    return env


def norm_label(s: str) -> str:
    s = (s or "").lower().strip()
    s = s.replace("\n", " ")
    s = re.sub(r"[^\w\s]", " ", s)
    s = re.sub(r"\s+", " ", s)
    return s.strip()


def parse_number(val: Any) -> float | None:
    if val is None:
        return None
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val).strip()
    if not s or s in ("-", "—", "N/A", "#DIV/0!", "#REF!", ""):
        return None
    neg = s.startswith("(") and s.endswith(")")
    if neg:
        s = s[1:-1]
    s = s.replace(",", "").replace(" ", "")
    try:
        n = float(s)
        return -n if neg else n
    except ValueError:
        return None


def get_fs_year_columns(ws, header_row: int = 4) -> dict[int, int]:
    mapping: dict[int, int] = {}
    for col in range(1, ws.max_column + 1):
        v = ws.cell(header_row, col).value
        if v is None:
            continue
        s = str(v).strip()
        m = re.match(r"^(\d{4})", s)
        if m:
            mapping[int(m.group(1))] = col
    return mapping


def _is_per_share_label(label: str) -> bool:
    nl = norm_label(label)
    return "per share" in nl or "eps" in nl or "earnings per" in nl


def _is_share_count_key(key: str) -> bool:
    """True for weighted-average / number-of-shares rows (not the EPS amount)."""
    nk = norm_label(key)
    return (
        "weighted average" in nk
        or "number of ordinary shares" in nk
        or "number of shares" in nk
        or nk.startswith("total number of shares")
    )


def _lookup_index_value(
    index: dict[str, float],
    template_label: str,
    patterns: list[str],
) -> float | None:
    """Match label→value preferring exact then longest containment; skip share-count rows for per-share labels."""
    per_share = _is_per_share_label(template_label)
    best_key = ""
    best_val: float | None = None

    def _consider(k: str, v: float) -> None:
        nonlocal best_key, best_val
        if per_share and _is_share_count_key(k):
            return
        if per_share and abs(float(v)) >= 1_000:
            # Per-share amounts are rarely >= 1000; huge hits are share counts / profit totals.
            return
        if len(k) > len(best_key):
            best_key = k
            best_val = v

    for pat in patterns:
        np = norm_label(pat)
        if not np:
            continue
        if np in index:
            v = index[np]
            if not (per_share and (_is_share_count_key(np) or abs(float(v)) >= 1_000)):
                return v
        for k, v in index.items():
            if np in k or k in np:
                _consider(k, v)
    if best_val is not None:
        return best_val

    nt = norm_label(template_label)
    if nt in index:
        v = index[nt]
        if not (per_share and (_is_share_count_key(nt) or abs(float(v)) >= 1_000)):
            return v
    for k, v in index.items():
        if nt in k or k in nt:
            _consider(k, v)
    return best_val


def find_value_col(doc: dict, year: int, entity_column: str = "group") -> int | None:
    """Pick the data column for `year`, preferring GROUP (or BANK) when headers allow."""
    headers = doc.get("header_rows") or []
    rows = doc.get("rows") or []
    if not rows:
        return None

    body_rows = [
        [str(c) for c in (row.get("cells") if isinstance(row, dict) else row or [])]
        for row in rows
    ]
    try:
        from comb_note_extractor import (
            find_annual_bank_body_col,
            find_annual_group_body_col,
            row_entity_year_col,
        )

        if entity_column.lower() == "bank":
            col = find_annual_bank_body_col(headers, body_rows, year)
        else:
            col = find_annual_group_body_col(headers, body_rows, year, entity_column)

        # Prefer per-row GROUP/BANK year columns (skips note + page no).
        sample_cols: list[int] = []
        for row in body_rows[:50]:
            if not (row[0] or "").strip():
                continue
            rc = row_entity_year_col(row, headers, year, entity_column)
            if rc is not None:
                sample_cols.append(rc)
        if sample_cols:
            from collections import Counter

            col = Counter(sample_cols).most_common(1)[0][0]
        elif col is not None:
            pass
        else:
            col = None
        if col is not None:
            return col
    except Exception:
        pass

    year_s = str(year)
    max_col = max(len(r.get("cells") or []) for r in rows)

    exact_cols = {
        ci
        for ci in range(1, max_col)
        for hrow in headers
        if ci < len(hrow) and str(hrow[ci]).strip() == year_s
    }
    search_cols = sorted(exact_cols) if exact_cols else list(range(1, max_col))

    best_col: int | None = None
    best_score = -1.0
    for ci in search_cols:
        magnitudes: list[float] = []
        big_count = 0
        page_like = 0
        for row in rows[:40]:
            cells = row.get("cells") if isinstance(row, dict) else []
            if ci >= len(cells):
                continue
            raw = str(cells[ci]).strip()
            v = parse_number(cells[ci])
            if v is None:
                continue
            av = abs(v)
            if av < 1_000 and "," not in raw and not raw.startswith("("):
                page_like += 1
                continue
            magnitudes.append(av)
            if av >= 100_000:
                big_count += 1

        if not magnitudes:
            continue

        median = sorted(magnitudes)[len(magnitudes) // 2]
        score = float(big_count * 10)
        if median >= 100_000:
            score += 20
        elif median >= 1_000:
            score += 8
        elif median < 500:
            score += 1  # EPS / ratios
        if page_like > big_count:
            score *= 0.1

        if score > best_score:
            best_score = score
            best_col = ci

    return best_col


class DataExtractor:
    DOC_PRIORITY = {
        ("annual", "income_statement"): 10,
        ("quarterly", "income_statement"): 8,
        ("annual", "sofp"): 9,
        ("annual", "cash_flows"): 7,
        ("annual", "ten_year_summary"): 5,
    }

    def __init__(self, db, company_slug: str):
        self.db = db
        self.company_slug = company_slug
        self._year_cache: dict[int, dict[str, float]] = {}

    _ANCHOR_LABELS = frozenset(
        {
            "gross income",
            "total assets",
            "total liabilities",
            "profit before income tax",
            "net cash flows from operating activities",
        }
    )

    def _doc_quality(self, doc: dict, year: int) -> float:
        col = find_value_col(doc, year)
        if col is None:
            return -1
        rows = doc.get("rows") or []
        big = 0
        anchor_hits = 0
        for row in rows[:60]:
            cells = row.get("cells") if isinstance(row, dict) else []
            label = str(cells[0] or "").strip().lower() if cells else ""
            if label in self._ANCHOR_LABELS:
                anchor_hits += 1
            if col < len(cells):
                v = parse_number(cells[col])
                if v is not None and abs(v) >= 100_000:
                    big += 1
        return big + len(rows) * 0.01 + anchor_hits * 50

    def _docs_for_year(self, year: int) -> list[dict]:
        docs: list[dict] = []
        for stmt in ("income_statement", "sofp", "cash_flows", "ten_year_summary"):
            docs.extend(
                list(
                    self.db.financial_tables.find(
                        {
                            "company_slug": self.company_slug,
                            "report_type": "annual",
                            "year": year,
                            "statement_key": stmt,
                        }
                    )
                )
            )
        if not any(d.get("statement_key") == "income_statement" for d in docs):
            docs.extend(
                list(
                    self.db.financial_tables.find(
                        {
                            "company_slug": self.company_slug,
                            "report_type": "quarterly",
                            "year": year,
                            "quarter": "Q4",
                            "statement_key": {"$regex": "^income_statement"},
                        }
                    )
                )
            )

        # Keep only the best table per statement_key (avoids bad duplicate extractions).
        best_by_key: dict[str, dict] = {}
        best_score: dict[str, float] = {}
        for doc in docs:
            sk = doc.get("statement_key") or ""
            q = self._doc_quality(doc, year)
            if sk not in best_score or q > best_score[sk]:
                best_score[sk] = q
                best_by_key[sk] = doc
        docs = list(best_by_key.values())

        def prio(d: dict) -> int:
            rt = d.get("report_type", "annual")
            sk = d.get("statement_key", "")
            for (r, s), p in self.DOC_PRIORITY.items():
                if rt == r and sk.startswith(s):
                    return p
            return 1

        return sorted(docs, key=prio, reverse=True)

    def _ingest_doc(self, index: dict[str, float], doc: dict, year: int) -> None:
        stmt = doc.get("statement_key") or ""
        col = find_value_col(doc, year)
        if col is None:
            return

        is_ten_year = stmt == "ten_year_summary"
        unit_scale = 1000.0 if is_ten_year else 1.0

        body_rows: list[list[str]] = []
        for row in doc.get("rows") or []:
            cells = row.get("cells") if isinstance(row, dict) else None
            if cells:
                body_rows.append([str(c) for c in cells])

        headers = doc.get("header_rows") or []
        from comb_note_extractor import index_label_values_from_rows

        for nl, scaled in index_label_values_from_rows(
            body_rows,
            col,
            header_rows=headers,
            year=year,
            unit_scale=unit_scale,
        ).items():
            if nl and nl not in index:
                index[nl] = scaled

    def index_for_year(self, year: int) -> dict[str, float]:
        if year not in self._year_cache:
            index: dict[str, float] = {}
            for doc in self._docs_for_year(year):
                self._ingest_doc(index, doc, year)
            self._year_cache[year] = index
        return self._year_cache[year]

    def lookup(self, year: int, template_label: str) -> float | None:
        index = self.index_for_year(year)
        patterns = patterns_for_label(
            template_label, LABEL_ALIASES, "fs", default_to_label=True
        )
        return _lookup_index_value(index, template_label, patterns)

def strip_forecast_columns(ws, first_extra_col: int = 12) -> None:
    """Remove forecast / interim columns from the FS sheet."""
    if ws.max_column >= first_extra_col:
        ws.delete_cols(first_extra_col, ws.max_column - first_extra_col + 1)


def clear_driver_formulas(ws, year_start_col: int = 3, year_end_col: int = 11) -> None:
    for row in ws.iter_rows(
        min_row=5,
        max_row=ws.max_row,
        min_col=year_start_col,
        max_col=year_end_col,
    ):
        for cell in row:
            value = cell.value
            if isinstance(value, str) and value.startswith("=") and "drivers" in value.lower():
                cell.value = None


def export_historical_fs_workbook(
    company_slug: str,
    output_path: Path,
    *,
    company_name: str | None = None,
    ticker: str | None = None,
    years: list[int] | None = None,
) -> dict[str, Any]:
    """Build Cover + FS workbook (2017-2025 actuals, no forecast columns)."""
    if not TEMPLATE.exists():
        raise FileNotFoundError(f"Template not found: {TEMPLATE}")

    env = load_env()
    uri = env.get("MONGO_URI", "mongodb://localhost:27017")
    db_name = env.get("MONGO_DB_NAME", "Research_Project")
    client = MongoClient(uri, serverSelectionTimeoutMS=5000)
    db = client[db_name]
    year_list = years or resolve_export_years(
        db, company_slug, report_type="annual", sheet="FS"
    )

    display_name = (company_name or "").strip()
    if not display_name:
        registry = db.companies.find_one({"slug": company_slug}, {"name": 1})
        if registry and registry.get("name"):
            display_name = str(registry["name"])
        else:
            sample = db.financial_tables.find_one(
                {"company_slug": company_slug},
                {"company_name": 1},
            )
            display_name = (
                str(sample.get("company_name"))
                if sample and sample.get("company_name")
                else company_slug.replace("_", " ")
            )

    wb = openpyxl.load_workbook(TEMPLATE)
    for sheet_name in list(wb.sheetnames):
        if sheet_name not in {"Cover", "FS"}:
            del wb[sheet_name]

    cover = wb["Cover"]
    cover.cell(3, 2, display_name)
    if ticker and ticker.strip():
        cover.cell(4, 2, ticker.strip())

    ws = wb["FS"]
    strip_forecast_columns(ws)
    for idx, year in enumerate(year_list):
        ws.cell(4, 3 + idx, year)

    extractor = DataExtractor(db, company_slug)
    template_rows = extract_template_labels(ws)
    filled = 0
    missing: list[dict[str, Any]] = []

    for row_num, label in template_rows:
        for year in year_list:
            col = 3 + year_list.index(year)
            val = extractor.lookup(year, label)
            if val is not None:
                ws.cell(row_num, col, val)
                filled += 1
            else:
                ws.cell(row_num, col, EMPTY_CELL)
                missing.append({"year": year, "row": row_num, "label": label})

    clear_driver_formulas(ws)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    client.close()

    return {
        "ok": True,
        "company_slug": company_slug,
        "company_name": display_name,
        "years": year_list,
        "output_file": str(output_path),
        "cells_filled": filled,
        "cells_missing": len(missing),
    }


def extract_template_rows(ws) -> list[dict[str, Any]]:
    """All FS rows for preview/export (sections + line items)."""
    rows: list[dict[str, Any]] = []
    for r in range(5, ws.max_row + 1):
        label = ws.cell(r, LABEL_COL).value
        if label is None:
            continue
        s = str(label).strip()
        if not s or s.startswith("For the year"):
            continue
        if s in SECTION_HEADERS:
            rows.append({"row": r, "label": s, "kind": "section"})
            continue
        if s in SUBSECTION_HEADERS:
            rows.append({"row": r, "label": s, "kind": "subsection"})
            continue
        if s in CHECK_LABELS:
            rows.append({"row": r, "label": s, "kind": "check"})
            continue
        if s in SKIP_LABELS:
            continue
        rows.append({"row": r, "label": s, "kind": "data"})
    return rows


def extract_template_labels(ws) -> list[tuple[int, str]]:
    rows: list[tuple[int, str]] = []
    for r in range(5, ws.max_row + 1):
        label = ws.cell(r, LABEL_COL).value
        if label is None:
            continue
        s = str(label).strip()
        if not s or s in SKIP_LABELS:
            continue
        if s.startswith("For the year"):
            continue
        rows.append((r, s))
    return rows


def _ui_note_tables_index(
    db,
    company_slug: str,
    years: list[int],
) -> dict[tuple[int, str], list[dict[str, Any]]]:
    """Preload the small OpenAI-extracted note-table UI pilot."""
    index: dict[tuple[int, str], list[dict[str, Any]]] = {}
    cursor = db.financial_tables.find(
        {
            "company_slug": company_slug,
            "year": {"$in": years},
            "report_type": "annual",
            "ui_extracted_tables.0": {"$exists": True},
        },
        {
            "year": 1,
            "statement_key": 1,
            "ui_extracted_tables": 1,
            "ui_extraction_method": 1,
            "ui_extraction_model": 1,
        },
    )
    for doc in cursor:
        year = int(doc.get("year") or 0)
        statement_key = str(doc.get("statement_key") or "")
        tables = doc.get("ui_extracted_tables") or []
        if not year or not statement_key or not tables:
            continue
        enriched: list[dict[str, Any]] = []
        for table in tables:
            if not isinstance(table, dict):
                continue
            enriched.append(
                {
                    **table,
                    "extraction_method": doc.get("ui_extraction_method"),
                    "extraction_model": doc.get("ui_extraction_model"),
                }
            )
        if enriched:
            index[(year, statement_key)] = enriched
    return index


def _note_tables_for_sources(
    table_index: dict[tuple[int, str], list[dict[str, Any]]],
    sources: dict[str, dict[str, Any]],
) -> dict[str, list[dict[str, Any]]]:
    by_year: dict[str, list[dict[str, Any]]] = {}
    for year_key, source in sources.items():
        statement_key = str(source.get("statement_key") or "")
        try:
            year = int(year_key)
        except (TypeError, ValueError):
            continue
        tables = table_index.get((year, statement_key))
        if tables:
            by_year[year_key] = tables
    return by_year


def build_fs_preview_data(
    company_slug: str,
    *,
    years: list[int] | None = None,
) -> dict[str, Any]:
    """Return FS grid JSON for in-app preview (same data as Excel export)."""
    if not TEMPLATE.exists():
        raise FileNotFoundError(f"Template not found: {TEMPLATE}")

    env = load_env()
    uri = env.get("MONGO_URI", "mongodb://localhost:27017")
    db_name = env.get("MONGO_DB_NAME", "Research_Project")

    client = MongoClient(uri, serverSelectionTimeoutMS=5000)
    db = client[db_name]
    year_list = years or resolve_export_years(
        db, company_slug, report_type="annual", sheet="FS"
    )
    store = CombWorkbookStore(db, company_slug)
    fs_note_links = build_note_metadata(TEMPLATE)
    ui_note_tables = _ui_note_tables_index(db, company_slug, year_list)

    wb = openpyxl.load_workbook(TEMPLATE, data_only=False)
    ws = wb["FS"]
    period_label = str(ws.cell(3, LABEL_COL).value or "For the year ended December 31,")
    template_rows = extract_template_rows(ws)

    preview_rows: list[dict[str, Any]] = []
    filled = 0
    missing = 0

    for item in template_rows:
        if item["kind"] == "check":
            continue
        if item["kind"] != "data":
            preview_rows.append(
                {"label": item["label"], "kind": item["kind"], "values": {}}
            )
            continue

        label = str(item["label"])
        values, statuses = _row_year_maps(store, "FS", label, year_list)
        note_meta = fs_note_links.get(label)
        note_source_by_year: dict[str, dict[str, Any]] = {}
        row_has_notes = False
        for year in year_list:
            cell_doc = store.cell_meta(year, label, sheet="FS")
            if cell_doc:
                if cell_doc.get("note_source"):
                    note_source_by_year[str(year)] = cell_doc["note_source"]
                if cell_doc.get("has_notes") or cell_doc.get("note_ref"):
                    row_has_notes = True
            if values.get(str(year)) is not None:
                filled += 1
            elif statuses.get(str(year)) == STATUS_CONFIRMED_ABSENT:
                missing += 1
            else:
                missing += 1
        row_payload: dict[str, Any] = {
            "label": label,
            "kind": "data",
            "values": values,
            "statuses": statuses,
            "row": item.get("row"),
        }
        if row_has_notes or note_meta:
            row_payload["has_notes"] = True
            if note_meta:
                row_payload["drivers_row"] = note_meta.get("drivers_row")
            if note_source_by_year:
                row_payload["note_source_by_year"] = note_source_by_year
                note_tables_by_year = _note_tables_for_sources(
                    ui_note_tables, note_source_by_year
                )
                if note_tables_by_year:
                    row_payload["note_tables_by_year"] = note_tables_by_year
        preview_rows.append(row_payload)

    template_validation = apply_template_validation_to_rows(
        preview_rows, "FS", year_list, use_row_key=True
    )

    client.close()
    wb.close()

    return {
        "view": "fs",
        "company_slug": company_slug,
        "years": year_list,
        "unit": "LKR '000 except per share data",
        "period_label": period_label,
        "rows": preview_rows,
        "cells_filled": filled,
        "cells_missing": missing,
        "template_validation": template_validation,
    }


def build_notes_preview_data(
    company_slug: str,
    *,
    years: list[int] | None = None,
) -> dict[str, Any]:
    """Return Notes tab JSON: FS descriptions with per-year note breakdown slots."""
    if not TEMPLATE.exists():
        raise FileNotFoundError(f"Template not found: {TEMPLATE}")

    env = load_env()
    uri = env.get("MONGO_URI", "mongodb://localhost:27017")
    db_name = env.get("MONGO_DB_NAME", "Research_Project")

    client = MongoClient(uri, serverSelectionTimeoutMS=5000)
    db = client[db_name]
    year_list = years or resolve_export_years(
        db, company_slug, report_type="annual", sheet="FS"
    )
    store = CombWorkbookStore(db, company_slug)
    fs_note_links = build_note_metadata(TEMPLATE)
    ui_note_tables = _ui_note_tables_index(db, company_slug, year_list)

    wb = openpyxl.load_workbook(TEMPLATE, data_only=False)
    ws = wb["FS"]
    period_label = str(ws.cell(3, LABEL_COL).value or "For the year ended December 31,")
    template_rows = extract_template_rows(ws)

    preview_rows: list[dict[str, Any]] = []
    notes_extracted_years: set[int] = set()
    note_rows_with_data = 0

    for item in template_rows:
        if item["kind"] == "check":
            continue
        if item["kind"] not in ("data", "check"):
            preview_rows.append(
                {"label": item["label"], "kind": item["kind"], "values": {}}
            )
            continue

        label = str(item["label"])
        values, statuses = _row_year_maps(store, "FS", label, year_list)
        row_payload: dict[str, Any] = {
            "label": label,
            "kind": item["kind"],
            "values": values,
            "statuses": statuses,
            "row": item.get("row"),
        }

        note_meta = fs_note_links.get(label)
        note_source_by_year: dict[str, dict[str, Any]] = {}
        row_has_notes = False
        for year in year_list:
            yr_key = str(year)
            cell_doc = store.cell_meta(year, label, sheet="FS")
            if cell_doc:
                if cell_doc.get("note_source"):
                    note_source_by_year[yr_key] = cell_doc["note_source"]
                    notes_extracted_years.add(year)
                if cell_doc.get("has_notes") or cell_doc.get("note_ref"):
                    row_has_notes = True
        if row_has_notes or note_meta:
            row_payload["has_notes"] = True
            if note_meta:
                row_payload["drivers_row"] = note_meta.get("drivers_row")
            if note_source_by_year:
                row_payload["note_source_by_year"] = note_source_by_year
                note_tables_by_year = _note_tables_for_sources(
                    ui_note_tables, note_source_by_year
                )
                if note_tables_by_year:
                    row_payload["note_tables_by_year"] = note_tables_by_year
            if any(note_source_by_year.get(str(y)) for y in year_list):
                note_rows_with_data += 1

        preview_rows.append(row_payload)

    client.close()
    wb.close()

    return {
        "view": "notes",
        "company_slug": company_slug,
        "years": year_list,
        "unit": "LKR '000 except per share data",
        "period_label": period_label,
        "rows": preview_rows,
        "notes_extracted_years": sorted(notes_extracted_years, reverse=True),
        "note_line_items": sum(1 for r in preview_rows if r.get("has_notes")),
        "note_line_items_with_data": note_rows_with_data,
    }


def _format_quarter_header(value: Any) -> str:
    if value is None:
        return ""
    if hasattr(value, "strftime"):
        return value.strftime("%b %Y")
    return str(value).strip()


def _safe_div(num: float | None, den: float | None) -> float | None:
    if num is None or den is None or den == 0:
        return None
    return num / den


def _fs_values_by_row(extractor: DataExtractor, year: int) -> dict[int, float | None]:
    cache_attr = "_fs_row_cache"
    cache: dict[int, dict[int, float | None]] = getattr(extractor, cache_attr, {})
    if year in cache:
        return cache[year]
    wb = openpyxl.load_workbook(TEMPLATE, data_only=False)
    ws = wb["FS"]
    row_vals: dict[int, float | None] = {}
    for row_num, label in extract_template_labels(ws):
        row_vals[row_num] = extractor.lookup(year, label)
    wb.close()
    cache[year] = row_vals
    setattr(extractor, cache_attr, cache)
    return row_vals


def _drivers_values_by_row(
    extractor: "DriversExtractor", year: int
) -> dict[int, float | None]:
    cache_attr = "_drv_row_cache"
    cache: dict[int, dict[int, float | None]] = getattr(extractor, cache_attr, {})
    if year in cache:
        return cache[year]
    wb = openpyxl.load_workbook(TEMPLATE, data_only=False)
    ws = wb["Drivers"]
    row_vals: dict[int, float | None] = {}
    for row_num, label in extract_template_labels(ws):
        row_vals[row_num] = extractor.lookup(year, label)
    wb.close()
    cache[year] = row_vals
    setattr(extractor, cache_attr, cache)
    return row_vals


_BREAKUPS_CACHE: dict[int, dict[str, float]] = {}


def _breakups_index_for_year(year: int) -> dict[str, float]:
    if year in _BREAKUPS_CACHE:
        return _BREAKUPS_CACHE[year]
    index: dict[str, float] = {}
    if not TEMPLATE.exists():
        _BREAKUPS_CACHE[year] = index
        return index
    wb = openpyxl.load_workbook(TEMPLATE, data_only=True)
    if "more Breakups" not in wb.sheetnames:
        wb.close()
        _BREAKUPS_CACHE[year] = index
        return index
    ws = wb["more Breakups"]
    year_col: int | None = None
    for col in range(3, ws.max_column + 1):
        header = ws.cell(1, col).value
        if header is not None and int(header) == year:
            year_col = col
            break
    if year_col is not None:
        for row in ws.iter_rows(min_row=4, max_row=ws.max_row, min_col=2, max_col=2):
            label = row[0].value
            if not label:
                continue
            val = ws.cell(row[0].row, year_col).value
            if isinstance(val, (int, float)):
                index[norm_label(str(label))] = float(val)
            else:
                parsed = parse_number(val)
                if parsed is not None:
                    index[norm_label(str(label))] = parsed
    wb.close()
    _BREAKUPS_CACHE[year] = index
    return index


def compute_ratios_for_year(
    fs_extractor: DataExtractor,
    drivers_extractor: "DriversExtractor",
    year: int,
) -> dict[str, float | None]:
    """Mirror COMB Ratios sheet formulas for one year."""
    def fs(label: str) -> float | None:
        return fs_extractor.lookup(year, label)

    interest_income = fs("Interest income")
    interest_expense = fs("Less: Interest expense")
    net_interest = fs("Net interest income")
    gross_income = fs("Gross income")
    total_operating_income = fs("Total operating income")
    impairment = fs("Less: Impairment charges and other losses")
    total_operating_expenses = fs("Total operating expenses")
    profit = fs("Profit for the year")
    total_equity = fs("Total equity")
    if total_equity is None:
        total_assets_for_eq = fs("Total assets")
        total_liabilities = fs("Total liabilities")
        if total_assets_for_eq is not None and total_liabilities is not None:
            total_equity = total_assets_for_eq - total_liabilities
    total_assets = fs("Total assets")
    deposits = fs("Financial liabilities at amortised cost – due to depositors")
    cash = fs("Cash and cash equivalents")

    earning_assets = sum(
        v or 0
        for v in [
            fs("Derivative financial assets"),
            fs(
                "Financial assets recognised through profit or loss – measured at fair value"
            ),
            fs("Financial assets at amortised cost – loans and advances to banks"),
            fs(
                "Financial assets at amortised cost – loans and advances to other customers"
            ),
            fs(
                "Financial assets at amortised cost – debt and other financial instruments"
            ),
            fs(
                "Financial assets measured at fair value through other comprehensive income"
            ),
        ]
    )
    bearing_funds = sum(
        v or 0
        for v in [
            fs("Due to banks"),
            fs("Derivative financial liabilities"),
            fs("Securities sold under repurchase agreements"),
            fs("Financial liabilities at amortised cost – due to depositors"),
            fs("Financial liabilities at amortised cost – other borrowings"),
        ]
    )

    dividend_per_share = fs("Diviend per share")
    if dividend_per_share is None and year in drivers_extractor.MEMORANDUM_BY_YEAR:
        dividend_per_share = drivers_extractor.MEMORANDUM_BY_YEAR[year].get(
            "dividend per share"
        )
    eps = fs("Basic earnings per ordinary share (Rs.)") or fs(
        "Diluted earnings per ordinary share (Rs.)"
    )
    employees = fs("Number of employees") or drivers_extractor.MEMORANDUM_BY_YEAR.get(
        year, {}
    ).get("number of employees")
    service_centres = fs("Number of customer service centres") or (
        drivers_extractor.MEMORANDUM_BY_YEAR.get(year, {}).get(
            "number of customer service centres"
        )
    )

    gross_loans = drivers_extractor.lookup(year, "Gross loans and advances") or fs(
        "Financial assets at amortised cost – loans and advances to other customers"
    )
    provision = drivers_extractor.lookup(year, "Less: Provision for impairment") or fs(
        "Less: Impairment charges and other losses"
    )

    avg_yield = _safe_div(interest_income, earning_assets or None)
    avg_cost = _safe_div(interest_expense, bearing_funds or None)
    nim = _safe_div(net_interest, earning_assets or None)
    spread = (
        (avg_yield - avg_cost)
        if avg_yield is not None and avg_cost is not None
        else None
    )

    return {
        "Avg. yield on interest earning assets": avg_yield,
        "Avg. cost on interest bearing funds": avg_cost,
        "Net interest margin": nim,
        "Interest rate spread": spread,
        "Cost to income (operating)": _safe_div(
            total_operating_expenses, total_operating_income
        ),
        "Cost to income (with impairments)": _safe_div(
            (total_operating_expenses or 0) + (impairment or 0),
            gross_income,
        )
        if gross_income
        else None,
        "Impairments/ Gross loans": _safe_div(provision, gross_loans),
        "Net margin": _safe_div(profit, gross_income),
        "ROE": _safe_div(profit, total_equity),
        "ROA": _safe_div(profit, total_assets),
        "Equity/ total assets": _safe_div(total_equity, total_assets),
        "Deposits / total assets": _safe_div(deposits, total_assets),
        "Cash / total assets": _safe_div(cash, total_assets),
        "Gross income per employee": _safe_div(gross_income, employees),
        "Dividend payout": _safe_div(dividend_per_share, eps),
        "Gross income per service centre": _safe_div(gross_income, service_centres),
        "Net interest spread": spread,
    }


class DriversExtractor(DataExtractor):
    """Broader annual lookup for Drivers sheet (all statement tables + breakups)."""

    MEMORANDUM_BY_YEAR: dict[int, dict[str, float]] = {
        2022: {
            "number of employees": 5121.0,
            "number of customer service centres": 289.0,
            "dividend per share": 4.5,
        },
    }
    MEMORANDUM_2022 = MEMORANDUM_BY_YEAR[2022]

    def _docs_for_year(self, year: int) -> list[dict]:
        """All annual tables (including notes) — best doc per statement_key."""
        docs = list(
            self.db.financial_tables.find(
                {
                    "company_slug": self.company_slug,
                    "report_type": "annual",
                    "year": year,
                }
            )
        )
        best_by_key: dict[str, dict] = {}
        best_score: dict[str, float] = {}
        for doc in docs:
            sk = doc.get("statement_key") or ""
            q = self._doc_quality(doc, year)
            if sk not in best_score or q > best_score[sk]:
                best_score[sk] = q
                best_by_key[sk] = doc
        return sorted(
            best_by_key.values(),
            key=lambda d: len(d.get("rows") or []),
            reverse=True,
        )

    def lookup(self, year: int, template_label: str) -> float | None:
        patterns = patterns_for_label(
            template_label, DRIVERS_LABEL_ALIASES, "drivers", default_to_label=False
        )
        index = self.index_for_year(year)
        for pat in patterns:
            np = norm_label(pat)
            if np in index:
                return index[np]
            for k, v in index.items():
                if np in k or k in np:
                    return v
        val = super().lookup(year, template_label)
        if val is not None:
            return val
        if year in self.MEMORANDUM_BY_YEAR:
            nt = norm_label(template_label)
            for k, v in self.MEMORANDUM_BY_YEAR[year].items():
                if k in nt or nt in k:
                    return v
        return _lookup_breakups_year(template_label, year)


def _lookup_breakups_year(template_label: str, year: int) -> float | None:
    """Fallback: more Breakups sheet holds note-level breakouts keyed by year column."""
    index = _breakups_index_for_year(year)
    nt = norm_label(template_label)
    if nt in index:
        return index[nt]
    for k, v in index.items():
        if nt in k or k in nt:
            return v
    return None


def extract_ratios_template_labels() -> list[str]:
    wb = openpyxl.load_workbook(TEMPLATE, data_only=False)
    ws = wb["Ratios"]
    labels: list[str] = []
    for r in range(3, ws.max_row + 1):
        label = ws.cell(r, 2).value
        if label and str(label).strip() and str(label).strip().lower() != "none":
            labels.append(str(label).strip())
    wb.close()
    return labels


def build_ratios_preview_data(
    company_slug: str,
    *,
    years: list[int] | None = None,
) -> dict[str, Any]:
    if not TEMPLATE.exists():
        raise FileNotFoundError(f"Template not found: {TEMPLATE}")

    env = load_env()
    client = MongoClient(env.get("MONGO_URI", "mongodb://localhost:27017"), serverSelectionTimeoutMS=5000)
    db = client[env.get("MONGO_DB_NAME", "Research_Project")]
    year_list = years or resolve_export_years(
        db, company_slug, report_type="annual", sheet="Ratios"
    )

    store = CombWorkbookStore(db, company_slug)
    ratio_labels = extract_ratios_template_labels()

    preview_rows: list[dict[str, Any]] = []
    filled = missing = 0
    for label in ratio_labels:
        values, statuses = _row_year_maps(store, "Ratios", label, year_list)
        for yk, val in values.items():
            if val is not None:
                filled += 1
            else:
                missing += 1
        preview_rows.append(
            {"label": label, "kind": "data", "values": values, "statuses": statuses}
        )

    template_validation = apply_template_validation_to_rows(
        preview_rows, "Ratios", year_list, is_ratio=True
    )

    client.close()
    return {
        "view": "ratios",
        "company_slug": company_slug,
        "years": year_list,
        "unit": "Ratios (%) except per-employee / per-centre figures",
        "period_label": "For the year ended December 31,",
        "rows": preview_rows,
        "cells_filled": filled,
        "cells_missing": missing,
        "value_format": "ratio",
        "template_validation": template_validation,
    }


def build_drivers_preview_data(
    company_slug: str,
    *,
    years: list[int] | None = None,
) -> dict[str, Any]:
    if not TEMPLATE.exists():
        raise FileNotFoundError(f"Template not found: {TEMPLATE}")

    env = load_env()
    client = MongoClient(env.get("MONGO_URI", "mongodb://localhost:27017"), serverSelectionTimeoutMS=5000)
    db = client[env.get("MONGO_DB_NAME", "Research_Project")]
    year_list = years or resolve_export_years(
        db, company_slug, report_type="annual", sheet="Drivers"
    )

    store = CombWorkbookStore(db, company_slug)
    wb = openpyxl.load_workbook(TEMPLATE, data_only=False)
    ws = wb["Drivers"]
    period_label = str(ws.cell(3, LABEL_COL).value or "For the year ended December 31,")
    template_rows = extract_template_rows(ws)
    wb.close()

    preview_rows: list[dict[str, Any]] = []
    filled = missing = 0
    for item in template_rows:
        if item["kind"] != "data":
            preview_rows.append(
                {"label": item["label"], "kind": item["kind"], "values": {}}
            )
            continue
        label = str(item["label"])
        row_num = int(item.get("row") or 0)
        values, statuses = _row_year_maps(
            store, "Drivers", label, year_list, drivers_row=row_num or None
        )
        for val in values.values():
            if val is not None:
                filled += 1
            else:
                missing += 1
        preview_rows.append(
            {
                "label": label,
                "kind": "data",
                "values": values,
                "statuses": statuses,
                "row": row_num,
            }
        )

    template_validation = apply_template_validation_to_rows(
        preview_rows, "Drivers", year_list, use_row_key=True
    )

    client.close()
    return {
        "view": "drivers",
        "company_slug": company_slug,
        "years": year_list,
        "unit": "LKR '000 except per share data",
        "period_label": period_label,
        "rows": preview_rows,
        "cells_filled": filled,
        "cells_missing": missing,
        "template_validation": template_validation,
    }


def _find_quarterly_group_col(doc: dict, year: int, entity_column: str = "group") -> int | None:
    """GROUP body column under 'For the quarter ended' (never six-month / Change %)."""
    from comb_note_extractor import find_quarter_ended_body_col

    headers = doc.get("header_rows") or []
    body_rows: list[list[str]] = []
    for row in doc.get("rows") or []:
        cells = row.get("cells") if isinstance(row, dict) else row
        if cells:
            body_rows.append([str(c) for c in cells])
    if headers:
        return find_quarter_ended_body_col(headers, body_rows, year, entity_column)
    return None


class QuarterlyExtractor:
    QUARTERLY_LABEL_ALIASES: dict[str, list[str]] = {
        "Gross income": ["gross income"],
        "Interest income": ["interest income"],
        "Less : Interest expense": [
            "less interest expense",
            "less : interest expense",
        ],
        "Net interest income": ["net interest income"],
        "Fee and commission income": ["fee and commission income"],
        "Less: Fee and commission expense": [
            "less fee and commission expense",
            "less: fee and commission expense",
        ],
        "Net fee and commission income": ["net fee and commission income"],
        "Net gains/(losses) from trading": ["net gains/(losses) from trading"],
        "Net gains/(losses) from derecognition of financial assets": [
            "net gains/(losses) from derecognition of financial assets",
        ],
        "Net other operating income": ["net other operating income"],
        "Total operating income": ["total operating income"],
        "Less : Impairment charges and other losses": [
            "less impairment charges and other losses",
            "less : impairment charges and other losses",
        ],
        "Net operating income": ["net operating income"],
        "Personnel expenses": ["personnel expenses"],
        "Depreciation and amortisation": ["depreciation and amortisation"],
        "Other operating expenses": ["other operating expenses"],
        "Operating profit before Taxes on financial services": [
            "operating profit before taxes on financial services",
        ],
        "Operating profit after Taxes on financial services": [
            "operating profit after taxes on financial services",
        ],
        "Add/(less): Share of profit/(loss) of associate, net of tax": [
            "add/(less): share of profit/(loss) of associate, net of tax",
            "share of profit/(loss) of associate",
        ],
        "Less: Taxes on financial services": ["less taxes on financial services"],
        "Profit before income tax": [
            "profit before income tax",
            "profit/(loss) before income tax",
            "profit before tax",
        ],
        "Less : Income tax expense / (reversal)": [
            "less income tax expense",
            "less : income tax expense",
        ],
        "Profit for the period": ["profit for the period", "profit for the year"],
        "Equity holders of the Bank": ["equity holders of the bank"],
        "Non-controlling interest": ["non-controlling interest"],
        "Earnings per share": ["earnings per share", "basic earnings per share"],
    }

    def __init__(self, db, company_slug: str, year: int):
        self.db = db
        self.company_slug = company_slug
        self.year = year
        self._cache: dict[str, dict[str, float]] = {}
        self._lookup_cache: dict[tuple[str, str], float | None] = {}

    def clear_cache(self, quarter: str | None = None) -> None:
        if quarter is None:
            self._cache.clear()
            self._lookup_cache.clear()
        else:
            self._cache.pop(quarter, None)
            self._lookup_cache = {
                k: v for k, v in self._lookup_cache.items() if k[0] != quarter
            }

    def _best_quarterly_doc(self, quarter: str) -> dict | None:
        docs = list(
            self.db.financial_tables.find(
                {
                    "company_slug": self.company_slug,
                    "report_type": "quarterly",
                    "year": self.year,
                    "quarter": quarter,
                    "statement_key": {"$regex": "^income_statement"},
                }
            )
        )
        if not docs:
            return None
        return max(docs, key=lambda d: len(d.get("rows") or []))

    def _index_quarter(self, quarter: str) -> dict[str, float]:
        if quarter in self._cache:
            return self._cache[quarter]
        index: dict[str, float] = {}
        doc = self._best_quarterly_doc(quarter)
        if doc:
            col = _find_quarterly_group_col(doc, self.year, entity_column="group")
            if col is not None:
                for row in doc.get("rows") or []:
                    cells = row.get("cells") or []
                    if not cells or not cells[0]:
                        continue
                    label = str(cells[0]).strip()
                    if col >= len(cells):
                        continue
                    val = parse_number(cells[col])
                    if val is None:
                        continue
                    nl = norm_label(label)
                    if nl and nl not in index:
                        index[nl] = val
        self._cache[quarter] = index
        return index

    def lookup(self, quarter: str, template_label: str) -> float | None:
        cache_key = (quarter, template_label)
        if cache_key in self._lookup_cache:
            return self._lookup_cache[cache_key]
        from comb_quarterly_validate import aggressive_quarterly_lookup

        val = aggressive_quarterly_lookup(self, quarter, template_label)
        self._lookup_cache[cache_key] = val
        return val


def build_quarterly_preview_data(
    company_slug: str = COMMERCIAL_BANK_SLUG,
    *,
    years: list[int] | None = None,
) -> dict[str, Any]:
    """Quarterly P&L from MongoDB (Commercial Bank pilot), extending by stored quarters."""
    if not TEMPLATE.exists():
        raise FileNotFoundError(f"Template not found: {TEMPLATE}")

    env = load_env()
    client = MongoClient(env.get("MONGO_URI", "mongodb://localhost:27017"), serverSelectionTimeoutMS=5000)
    db = client[env.get("MONGO_DB_NAME", "Research_Project")]
    year_list = years or resolve_export_years(
        db, company_slug, report_type="quarterly", sheet="Quarterly"
    )

    wb = openpyxl.load_workbook(TEMPLATE, data_only=False)
    ws = wb["Quarterly"]
    period_label = str(ws.cell(1, 3).value or "For the quarter ended")
    template_rows: list[dict[str, Any]] = []
    for r in range(3, ws.max_row + 1):
        label = ws.cell(r, 2).value
        if label and str(label).strip():
            label_s = str(label).strip()
            template_rows.append({"label": label_s, "row": r})

    columns = [
        {"key": col["key"], "label": col["label"]}
        for col in quarterly_column_defs(year_list)
    ]

    store = CombWorkbookStore(db, company_slug)
    rows: list[dict[str, Any]] = []
    filled = missing = 0
    for item in template_rows:
        label = item["label"]
        template_row = int(item["row"])
        kind = "subsection" if label.lower() == "profit attributable to:" else "data"
        values: dict[str, float | None] = {}
        statuses: dict[str, str] = {}
        for col_def in quarterly_column_defs(year_list):
            target_year = int(col_def["year"])
            quarter = str(col_def["quarter"])
            raw = store.lookup(
                target_year,
                label,
                sheet="Quarterly",
                report_type="quarterly",
                quarter=quarter,
            )
            val = raw
            values[col_def["key"]] = val
            statuses[col_def["key"]] = store.lookup_status(
                target_year,
                label,
                sheet="Quarterly",
                report_type="quarterly",
                quarter=quarter,
            )
            if val is not None:
                filled += 1
            else:
                missing += 1
        rows.append(
            {
                "label": label,
                "kind": kind,
                "template_row": template_row,
                "values": values,
                "statuses": statuses,
            }
        )

    wb.close()

    column_keys = [c["key"] for c in columns]
    template_validation = apply_quarterly_template_validation(rows, column_keys)

    client.close()
    return {
        "view": "quarterly",
        "company_slug": company_slug,
        "company_name": "Commercial Bank of Ceylon PLC",
        "unit": "LKR '000 except per share data",
        "period_label": period_label,
        "columns": columns,
        "rows": rows,
        "cells_filled": filled,
        "cells_missing": missing,
        "years": year_list,
        "template_validation": template_validation,
    }


def _write_year_grid(
    ws,
    template_rows: list[dict[str, Any]],
    year_list: list[int],
    lookup_fn,
    status_fn=None,
) -> tuple[int, int]:
    filled = missing = 0
    for idx, year in enumerate(year_list):
        ws.cell(4, 3 + idx, year)
    for item in template_rows:
        if item["kind"] != "data":
            continue
        row_num = item["row"]
        label = str(item["label"])
        for y_idx, year in enumerate(year_list):
            val = lookup_fn(year, label, row_num)
            cell = ws.cell(row_num, 3 + y_idx)
            status = status_fn(year, label, row_num) if status_fn else None
            if val is not None:
                cell.value = val
                filled += 1
            else:
                cell.value = EMPTY_CELL
                missing += 1
            if status == STATUS_CONFIRMED_ABSENT or status == STATUS_TEMPLATE_MISMATCH:
                cell.fill = RED_FILL
    return filled, missing


def export_comb_workbook(
    company_slug: str,
    output_path: Path,
    *,
    company_name: str | None = None,
    ticker: str | None = None,
    years: list[int] | None = None,
) -> dict[str, Any]:
    """Export Cover + FS + Drivers + Ratios + Quarterly for pilot year(s)."""
    if not TEMPLATE.exists():
        raise FileNotFoundError(f"Template not found: {TEMPLATE}")

    env = load_env()
    client = MongoClient(env.get("MONGO_URI", "mongodb://localhost:27017"), serverSelectionTimeoutMS=5000)
    db = client[env.get("MONGO_DB_NAME", "Research_Project")]
    annual_years = years or resolve_export_years(
        db, company_slug, report_type="annual"
    )
    quarterly_years = years or resolve_export_years(
        db, company_slug, report_type="quarterly", sheet="Quarterly"
    )

    display_name = (company_name or "").strip()
    if not display_name:
        sample = db.financial_tables.find_one(
            {"company_slug": company_slug}, {"company_name": 1}
        )
        display_name = (
            str(sample.get("company_name"))
            if sample and sample.get("company_name")
            else company_slug.replace("_", " ")
        )

    store = CombWorkbookStore(db, company_slug)

    wb = openpyxl.load_workbook(TEMPLATE)
    for sheet_name in list(wb.sheetnames):
        if sheet_name not in EXPORT_SHEETS:
            del wb[sheet_name]

    cover = wb["Cover"]
    cover.cell(3, 2, display_name)
    if ticker and ticker.strip():
        cover.cell(4, 2, ticker.strip())

    total_filled = total_missing = 0

    def fs_lookup(y: int, lbl: str, row_num: int | None = None) -> float | None:
        return display_value_for_year(y, store.lookup(y, lbl, sheet="FS"))

    def fs_status(y: int, lbl: str, row_num: int | None = None) -> str:
        return store.lookup_status(y, lbl, sheet="FS")

    def drv_lookup(y: int, lbl: str, row_num: int | None = None) -> float | None:
        return display_value_for_year(
            y, store.lookup(y, lbl, sheet="Drivers", drivers_row=row_num)
        )

    def drv_status(y: int, lbl: str, row_num: int | None = None) -> str:
        return store.lookup_status(y, lbl, sheet="Drivers", drivers_row=row_num)

    # FS
    fs_ws = wb["FS"]
    strip_forecast_columns(fs_ws)
    fs_rows = extract_template_rows(fs_ws)
    f, m = _write_year_grid(fs_ws, fs_rows, annual_years, fs_lookup, fs_status)
    total_filled += f
    total_missing += m
    clear_driver_formulas(fs_ws, 3, 3 + len(annual_years) - 1)

    # Drivers
    drv_ws = wb["Drivers"]
    strip_forecast_columns(drv_ws)
    drv_rows = extract_template_rows(drv_ws)
    f, m = _write_year_grid(drv_ws, drv_rows, annual_years, drv_lookup, drv_status)
    total_filled += f
    total_missing += m

    # Ratios
    ratios_ws = wb["Ratios"]
    strip_forecast_columns(ratios_ws)
    ratio_labels = extract_ratios_template_labels()
    for idx, year in enumerate(annual_years):
        ratios_ws.cell(4, 3 + idx, year)
    for r, label in enumerate(ratio_labels, start=3):
        ratios_ws.cell(r, 2, label)
        for y_idx, year in enumerate(annual_years):
            val = display_value_for_year(
                year, store.lookup(year, label, sheet="Ratios")
            )
            cell = ratios_ws.cell(r, 3 + y_idx)
            status = store.lookup_status(year, label, sheet="Ratios")
            if val is not None:
                cell.value = val
                total_filled += 1
            else:
                cell.value = EMPTY_CELL
                total_missing += 1
            if status == STATUS_CONFIRMED_ABSENT or status == STATUS_TEMPLATE_MISMATCH:
                cell.fill = RED_FILL

    # Quarterly — extends when newer quarter data is stored
    q_ws = wb["Quarterly"]
    q_cols = quarterly_column_defs(quarterly_years)
    q_labels: list[str] = []
    for r in range(3, q_ws.max_row + 1):
        label = q_ws.cell(r, 2).value
        if label and str(label).strip():
            q_labels.append(str(label).strip())
    for q_idx, col_def in enumerate(q_cols):
        q_ws.cell(2, 3 + q_idx, col_def["label"])
    for r, label in enumerate(q_labels, start=3):
        q_ws.cell(r, 2, label)
        for q_idx, col_def in enumerate(q_cols):
            val = store.lookup(
                int(col_def["year"]),
                label,
                sheet="Quarterly",
                report_type="quarterly",
                quarter=str(col_def["quarter"]),
            )
            cell = q_ws.cell(r, 3 + q_idx)
            status = store.lookup_status(
                int(col_def["year"]),
                label,
                sheet="Quarterly",
                report_type="quarterly",
                quarter=str(col_def["quarter"]),
            )
            if val is not None:
                cell.value = val
                total_filled += 1
            else:
                cell.value = EMPTY_CELL
                total_missing += 1
            if status == STATUS_CONFIRMED_ABSENT or status == STATUS_TEMPLATE_MISMATCH:
                cell.fill = RED_FILL
    extra_cols = q_ws.max_column - 2 - len(q_cols)
    if extra_cols > 0:
        q_ws.delete_cols(3 + len(q_cols), extra_cols)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    client.close()

    return {
        "ok": True,
        "company_slug": company_slug,
        "company_name": display_name,
        "years": annual_years,
        "quarterly_years": quarterly_years,
        "output_file": str(output_path),
        "cells_filled": total_filled,
        "cells_missing": total_missing,
        "sheets": list(EXPORT_SHEETS),
    }


def main() -> int:
    env = load_env()
    uri = env.get("MONGO_URI", "mongodb://localhost:27017")
    db_name = env.get("MONGO_DB_NAME", "Research_Project")

    if not TEMPLATE.exists():
        print(f"Template not found: {TEMPLATE}", file=sys.stderr)
        return 1

    client = MongoClient(uri, serverSelectionTimeoutMS=5000)
    db = client[db_name]
    extractor = DataExtractor(db, COMPANY_SLUG)

    # Coverage summary
    coverage = list(
        db.financial_tables.aggregate(
            [
                {"$match": {"company_slug": COMPANY_SLUG}},
                {
                    "$group": {
                        "_id": {
                            "year": "$year",
                            "report_type": "$report_type",
                            "quarter": "$quarter",
                        },
                        "cnt": {"$sum": 1},
                        "statements": {"$addToSet": "$statement_key"},
                    }
                },
                {"$sort": {"_id.year": 1}},
            ]
        )
    )

    wb = openpyxl.load_workbook(TEMPLATE)
    ws = wb["FS"]
    year_cols = get_fs_year_columns(ws)

    # Set year headers 2017-2025 in columns C-K (3-11)
    for i, year in enumerate(YEARS):
        ws.cell(4, 3 + i, year)

    template_rows = extract_template_labels(ws)
    missing: list[dict] = []
    filled = 0
    not_found_labels: set[str] = set()

    for row_num, label in template_rows:
        for year in YEARS:
            col = year_cols.get(year)
            if col is None and year in YEARS:
                col = 3 + YEARS.index(year)
            val = extractor.lookup(year, label)
            if val is not None:
                ws.cell(row_num, col, val)
                filled += 1
            else:
                missing.append({"year": year, "row": row_num, "label": label, "sheet": "FS"})
                not_found_labels.add(label)

    wb.save(OUTPUT)

    by_label: dict[str, list[int]] = {}
    for m in missing:
        by_label.setdefault(m["label"], []).append(m["year"])

    annual_years = sorted(
        {
            c["_id"]["year"]
            for c in coverage
            if c["_id"].get("report_type") == "annual" and c["_id"].get("year")
        }
    )

    report = {
        "company": COMPANY_SLUG,
        "years_requested": YEARS,
        "output_file": str(OUTPUT),
        "cells_filled_from_database": filled,
        "cells_still_missing": len(missing),
        "annual_report_years_in_db": annual_years,
        "gaps": {
            "2023_annual_income_statement": "income_statement"
            not in next(
                (
                    c["statements"]
                    for c in coverage
                    if c["_id"].get("year") == 2023
                    and c["_id"].get("report_type") == "annual"
                ),
                [],
            ),
        },
        "missing_by_label": {
            lbl: sorted(set(yrs)) for lbl, yrs in sorted(by_label.items())
        },
        "labels_never_found_in_any_year": sorted(not_found_labels),
        "annual_coverage": [
            {
                "year": c["_id"]["year"],
                "report_type": c["_id"]["report_type"],
                "quarter": c["_id"]["quarter"],
                "table_count": c["cnt"],
                "statements": sorted(c["statements"]),
            }
            for c in coverage
        ],
    }
    MISSING_REPORT.write_text(json.dumps(report, indent=2), encoding="utf-8")

    print(f"Generated: {OUTPUT}")
    print(f"Missing report: {MISSING_REPORT}")
    print(f"Filled {filled} cells; {len(missing)} still missing")
    print(f"Labels with gaps: {len(by_label)}")
    client.close()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
