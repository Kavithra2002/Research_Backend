"""
Build extraction manifest from COMB model - updated.xlsx.

The manifest is the contract for comb_workbook_data extraction.
"""
from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any

import openpyxl

from comb_workbook_store import (
    LABEL_COL,
    extract_fs_driver_links,
    extract_template_rows,
    resolve_template,
)

MANIFEST_CACHE = Path(__file__).resolve().parent.parent / "New_Updates" / "comb_manifest.json"


def _is_check_label(label: str) -> bool:
    return label.lower().strip() == "check" or "check" in label.lower() and len(label) < 12


def _is_skip_child(label: str) -> bool:
    low = label.lower()
    return low.startswith("as a %") or low.startswith("as a percent")


def build_drivers_groups(drv_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Group Drivers rows into parent + children blocks separated by Check rows."""
    groups: list[dict[str, Any]] = []
    segment: list[dict[str, Any]] = []

    def flush_segment() -> None:
        if not segment:
            return
        parent = segment[0]
        children = segment[1:]
        groups.append(
            {
                "parent_label": str(parent["label"]),
                "parent_row": int(parent["row"]),
                "children": [
                    {"label": str(c["label"]), "row": int(c["row"])} for c in children
                ],
            }
        )

    for item in drv_rows:
        if item["kind"] == "check" or _is_check_label(str(item["label"])):
            flush_segment()
            segment = []
            continue
        if item["kind"] != "data" or _is_skip_child(str(item["label"])):
            continue
        segment.append(item)

    flush_segment()
    return groups


def build_manifest(template_path: Path | None = None) -> dict[str, Any]:
    path = template_path or resolve_template()
    wb = openpyxl.load_workbook(path, data_only=False)
    fs_ws = wb["FS"]
    drv_ws = wb["Drivers"]
    q_ws = wb["Quarterly"]

    fs_rows = extract_template_rows(fs_ws)
    drv_rows = extract_template_rows(drv_ws)
    fs_links = extract_fs_driver_links(fs_ws)

    quarterly_rows: list[dict[str, Any]] = []
    for r in range(3, q_ws.max_row + 1):
        label = q_ws.cell(r, 2).value
        if label and str(label).strip():
            quarterly_rows.append({"label": str(label).strip(), "row": r})

    ratio_labels: list[str] = []
    if "Ratios" in wb.sheetnames:
        rws = wb["Ratios"]
        for r in range(3, rws.max_row + 1):
            label = rws.cell(r, 2).value
            if label and str(label).strip().lower() != "none":
                ratio_labels.append(str(label).strip())

    wb.close()

    manifest = {
        "template": str(path),
        "fs": {
            "rows": [
                {"label": i["label"], "row": i["row"], "kind": i["kind"]}
                for i in fs_rows
            ],
            "driver_links": fs_links,
        },
        "drivers": {
            "rows": [
                {"label": i["label"], "row": i["row"], "kind": i["kind"]}
                for i in drv_rows
            ],
            "groups": build_drivers_groups(drv_rows),
        },
        "quarterly": {
            "labels": [row["label"] for row in quarterly_rows],
            "rows": quarterly_rows,
        },
        "ratios": {"labels": ratio_labels},
        "entity_column": "group",
    }
    return manifest


def load_manifest(refresh: bool = False) -> dict[str, Any]:
    if not refresh and MANIFEST_CACHE.exists():
        try:
            return json.loads(MANIFEST_CACHE.read_text(encoding="utf-8"))
        except Exception:
            pass
    manifest = build_manifest()
    MANIFEST_CACHE.parent.mkdir(parents=True, exist_ok=True)
    MANIFEST_CACHE.write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    return manifest
