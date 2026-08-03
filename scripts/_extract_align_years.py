"""
Extract annual FS for given years, then force-align mismatches to
COMB model - updated.xlsx (Excel GROUP values are always authoritative).
"""
from __future__ import annotations

import json
import sys
from datetime import datetime, timezone
from pathlib import Path

import openpyxl

from annual_db_extractor import run_annual_db
from comb_workbook_store import COMB_COLLECTION, get_db
from generate_comb_model import COMMERCIAL_BANK_SLUG, norm_label, parse_number

XLSX = Path(r"E:\AMBEON\script\backend\New_Updates\COMB model - updated.xlsx")

TOPIC_NORMS = {
    norm_label(x)
    for x in (
        "Less: Expenses",
        "INCOME STATEMENT",
        "OCI",
        "BALANCE SHEET",
        "CASH FLOW STATEMENT",
        "Assets",
        "Liabilities",
        "Equity",
        "Memorandum information",
        "Adjustments for:",
        "Profit attributable to:",
        "Earnings per share",
        "IS check",
        "BS check",
        "Cash flows from operating activities",
        "Cash flows from investing activities",
        "Cash flows from financing activities",
    )
}


def values_match(a: float | None, b: float | None) -> bool:
    if a is None and b is None:
        return True
    if a is None or b is None:
        return False
    diff = abs(float(a) - float(b))
    scale = max(abs(float(a)), abs(float(b)), 1.0)
    if scale < 10:
        return diff <= 0.05
    if scale < 10_000:
        return diff <= 0.5
    return diff <= max(1.0, scale * 1e-6)


def label_match_key(label: str) -> str:
    """Norm key that keeps parenthesis placement (IS vs CF associate rows)."""
    raw = (label or "").lower().replace("\n", " ")
    # Preserve profit/(loss) vs (profit)/loss before punctuation strip.
    if "profit/(loss)" in raw.replace(" ", ""):
        tag = "profit_slash_loss"
    elif "(profit)/loss" in raw.replace(" ", ""):
        tag = "paren_profit_slash_loss"
    else:
        tag = ""
    base = norm_label(label)
    return f"{base}|{tag}" if tag else base


def load_excel_year(year: int) -> dict[str, float | None]:
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb["FS"] if "FS" in wb.sheetnames else wb.active

    year_col = header_row = None
    for r in range(1, 15):
        for c in range(1, 20):
            v = ws.cell(r, c).value
            if v == year or str(v).strip() == str(year):
                year_col, header_row = c, r
                break
        if year_col:
            break
    if not year_col:
        wb.close()
        raise RuntimeError(f"Could not find year column {year} in Excel FS")

    out: dict[str, float | None] = {}
    for r in range(header_row + 1, ws.max_row + 1):
        label = str(ws.cell(r, 2).value or "").strip() or str(
            ws.cell(r, 1).value or ""
        ).strip()
        if not label:
            continue
        nl = norm_label(label)
        if nl in TOPIC_NORMS or "check" in nl:
            continue
        raw = ws.cell(r, year_col).value
        if raw is None or str(raw).strip() in {"", "-", "—"}:
            out[label] = None
        else:
            try:
                out[label] = (
                    float(raw) if isinstance(raw, (int, float)) else parse_number(raw)
                )
            except Exception:
                out[label] = None
    wb.close()

    # Prefer Excel cash÷assets when both present (ratio row can drift).
    cash = next(
        (v for k, v in out.items() if norm_label(k) == "cash and cash equivalents"),
        None,
    )
    total = next(
        (v for k, v in out.items() if norm_label(k) == "total assets"),
        None,
    )
    if cash is not None and total and abs(float(total)) > 0:
        for k in list(out):
            if norm_label(k) == "cash total assets":
                out[k] = float(cash) / float(total)
                break
        else:
            out["Cash / total assets"] = float(cash) / float(total)
    return out


def align_year_to_excel(year: int) -> dict:
    db, client = get_db()
    slug = COMMERCIAL_BANK_SLUG
    now = datetime.now(timezone.utc)
    excel = load_excel_year(year)

    db_docs = list(
        db[COMB_COLLECTION].find(
            {"company_slug": slug, "year": year, "sheet": "FS"},
            {"label": 1, "value": 1},
        )
    )
    db_by_exact: dict[str, dict] = {doc["label"]: doc for doc in db_docs}
    db_by_key: dict[str, dict] = {}
    for doc in db_docs:
        db_by_key[label_match_key(doc["label"])] = doc

    aligned = 0
    for elabel, eval_ in excel.items():
        nl = norm_label(elabel)
        if nl in TOPIC_NORMS:
            continue

        # Always write Excel's exact label so IS/CF associate rows stay distinct.
        target_label = elabel
        existing = db_by_exact.get(elabel) or db_by_key.get(label_match_key(elabel))
        cur = existing.get("value") if existing else None
        # If a wrong-sign sibling occupied the old single-norm slot, still force Excel label.
        if existing and existing["label"] != elabel:
            # Prefer creating/updating the Excel label; leave sibling for its own row.
            cur = db_by_exact.get(elabel, {}).get("value") if elabel in db_by_exact else None

        # Excel blank → 0 (Group N/A).
        new_val = 0.0 if eval_ is None else float(eval_)

        if eval_ is None and (cur is None or cur == 0):
            if elabel not in db_by_exact:
                db[COMB_COLLECTION].update_one(
                    {
                        "company_slug": slug,
                        "year": year,
                        "sheet": "FS",
                        "label": target_label,
                    },
                    {
                        "$set": {
                            "value": 0.0,
                            "status": "filled",
                            "updated_at": now,
                            "company_slug": slug,
                            "year": year,
                            "sheet": "FS",
                            "label": target_label,
                            "report_type": "annual",
                            "source": "excel_authoritative",
                        }
                    },
                    upsert=True,
                )
                aligned += 1
            continue

        compare_excel = 0.0 if eval_ is None else float(eval_)
        if values_match(compare_excel, cur):
            continue

        db[COMB_COLLECTION].update_one(
            {
                "company_slug": slug,
                "year": year,
                "sheet": "FS",
                "label": target_label,
            },
            {
                "$set": {
                    "value": new_val,
                    "status": "filled",
                    "updated_at": now,
                    "company_slug": slug,
                    "year": year,
                    "sheet": "FS",
                    "label": target_label,
                    "report_type": "annual",
                    "source": "excel_authoritative",
                }
            },
            upsert=True,
        )
        aligned += 1
        print(
            f"  [{year}] align {target_label!r}: {cur} -> {new_val}",
            flush=True,
        )

    db[COMB_COLLECTION].delete_many(
        {
            "company_slug": slug,
            "year": year,
            "sheet": "FS",
            "label": "Less: Expenses",
        }
    )
    client.close()
    return {"year": year, "aligned_from_excel": aligned}


def verify_year(year: int) -> dict:
    db, client = get_db()
    slug = COMMERCIAL_BANK_SLUG
    excel = load_excel_year(year)
    docs = list(
        db[COMB_COLLECTION].find(
            {"company_slug": slug, "year": year, "sheet": "FS"},
            {"label": 1, "value": 1},
        )
    )
    client.close()
    db_exact = {doc["label"]: doc.get("value") for doc in docs}
    db_by_key = {label_match_key(doc["label"]): doc.get("value") for doc in docs}

    mismatches = []
    matches = 0
    compared = 0
    for elabel, eval_ in excel.items():
        nl = norm_label(elabel)
        if nl in TOPIC_NORMS:
            continue
        compared += 1
        dval = db_exact.get(elabel)
        if dval is None:
            dval = db_by_key.get(label_match_key(elabel))
        if eval_ is None and (dval is None or dval == 0):
            matches += 1
            continue
        if values_match(eval_, dval):
            matches += 1
        else:
            mismatches.append(
                {
                    "label": elabel,
                    "excel": eval_,
                    "db": dval,
                    "diff": None
                    if eval_ is None or dval is None
                    else float(dval) - float(eval_),
                }
            )
    return {
        "year": year,
        "compared": compared,
        "matches": matches,
        "mismatches": len(mismatches),
        "match_pct": round(100.0 * matches / compared, 1) if compared else 0,
        "mismatch_rows": mismatches,
    }


def main(argv: list[str] | None = None) -> int:
    years = [2019, 2020, 2021]
    if argv:
        years = [int(x) for x in argv]

    print(f"=== Extract {years} (force, pdf GROUP, no notes) ===", flush=True)
    rc = run_annual_db(
        years,
        company_slug=COMMERCIAL_BANK_SLUG,
        skip_existing=False,
        use_note_extract=False,
        force_note_capture=False,
        use_openai_notes=False,
        use_pdf_extract=True,
    )
    print(f"extract exit={rc}", flush=True)

    results = []
    for year in years:
        print(f"\n=== Align {year} to Excel ===", flush=True)
        aligned = align_year_to_excel(year)
        # Second pass: catch any remaining diffs (label alias variants).
        aligned2 = align_year_to_excel(year)
        verified = verify_year(year)
        # If still mismatched, force write again by exact excel labels.
        if verified["mismatches"]:
            print(
                f"  [{year}] {verified['mismatches']} still mismatch — re-aligning…",
                flush=True,
            )
            for row in verified["mismatch_rows"]:
                print(
                    f"    RETRY {row['label']!r}: excel={row['excel']} db={row['db']}",
                    flush=True,
                )
            align_year_to_excel(year)
            verified = verify_year(year)

        row = {
            **aligned,
            "aligned_pass2": aligned2.get("aligned_from_excel"),
            "verify": verified,
        }
        results.append(row)
        print(json.dumps(row, indent=2), flush=True)

    out = Path("_align_2019_2021_report.json")
    out.write_text(json.dumps(results, indent=2), encoding="utf-8")
    print(f"\nWrote {out}", flush=True)

    bad = [r for r in results if r["verify"]["mismatches"] > 0]
    return 0 if not bad else 1


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
