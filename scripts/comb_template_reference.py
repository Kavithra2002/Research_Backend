"""
Load reference values from COMB model - updated.xlsx for validation against extracted data.
"""
from __future__ import annotations

import re
from functools import lru_cache
from pathlib import Path
from typing import Any

import openpyxl

from comb_cell_status import PILOT_YEARS, STATUS_FILLED, STATUS_TEMPLATE_MISMATCH
from comb_workbook_store import LABEL_COL, resolve_template

LABEL_COL_REF = LABEL_COL

QUARTERLY_PILOT_QUARTERS = ["Q1", "Q2", "Q3", "Q4"]
QUARTERLY_PILOT_YEAR = 2022


def parse_number(val) -> float | None:
    from generate_comb_model import parse_number as _pn

    return _pn(val)


def extract_template_rows(ws):
    from generate_comb_model import extract_template_rows as _etr

    return _etr(ws)


def _year_columns(ws, header_row: int = 4) -> dict[int, int]:
    """Map calendar year -> column index for FS/Drivers-style sheets."""
    mapping: dict[int, int] = {}
    for col in range(3, ws.max_column + 1):
        v = ws.cell(header_row, col).value
        if v is None:
            continue
        s = str(v).strip()
        m = re.match(r"^(\d{4})", s)
        if m:
            mapping[int(m.group(1))] = col
    return mapping


def _ratios_year_columns(ws) -> dict[int, int]:
    """Ratios sheet stores years on row 2."""
    mapping: dict[int, int] = {}
    for col in range(3, ws.max_column + 1):
        v = ws.cell(2, col).value
        if v is None:
            continue
        s = str(v).strip()
        m = re.match(r"^(\d{4})", s)
        if m:
            mapping[int(m.group(1))] = col
    return mapping


def _quarterly_columns(ws) -> dict[str, int]:
    """Map q1_2022-style keys to column index for pilot quarter columns."""
    mapping: dict[str, int] = {}
    for q_idx, quarter in enumerate(QUARTERLY_PILOT_QUARTERS):
        col = 3 + q_idx
        if col > ws.max_column:
            break
        mapping[f"{quarter.lower()}_{QUARTERLY_PILOT_YEAR}"] = col
    return mapping


@lru_cache(maxsize=4)
def _load_workbook_data_only() -> openpyxl.Workbook:
    return openpyxl.load_workbook(resolve_template(), data_only=True)


def reference_by_row(
    sheet: str,
    *,
    years: list[int] | None = None,
) -> dict[tuple[int, int], float]:
    """(template_row, year) -> reference value."""
    wb = _load_workbook_data_only()
    if sheet not in wb.sheetnames:
        return {}
    ws = wb[sheet]
    year_cols = _year_columns(ws)
    target_years = years or PILOT_YEARS
    out: dict[tuple[int, int], float] = {}
    for year in target_years:
        col = year_cols.get(year)
        if col is None:
            continue
        for row in range(5, ws.max_row + 1):
            raw = ws.cell(row, col).value
            val = parse_number(raw)
            if val is not None:
                out[(row, year)] = val
    return out


def reference_by_label(
    sheet: str,
    *,
    years: list[int] | None = None,
) -> dict[tuple[str, int], float]:
    """(label, year) -> reference value."""
    wb = _load_workbook_data_only()
    if sheet not in wb.sheetnames:
        return {}
    ws = wb[sheet]
    year_cols = _year_columns(ws) if sheet != "Ratios" else _ratios_year_columns(ws)
    target_years = years or PILOT_YEARS
    out: dict[tuple[str, int], float] = {}
    for year in target_years:
        col = year_cols.get(year)
        if col is None:
            continue
        for row in range(3 if sheet == "Ratios" else 5, ws.max_row + 1):
            label = ws.cell(row, LABEL_COL_REF).value
            if not label:
                continue
            label_s = str(label).strip()
            raw = ws.cell(row, col).value
            val = parse_number(raw)
            if val is not None:
                out[(label_s, year)] = val
    return out


def reference_quarterly() -> dict[tuple[str, str], float]:
    """(label, column_key) -> reference value for pilot quarters."""
    wb = _load_workbook_data_only()
    if "Quarterly" not in wb.sheetnames:
        return {}
    ws = wb["Quarterly"]
    col_map = _quarterly_columns(ws)
    out: dict[tuple[str, str], float] = {}
    for col_key, col in col_map.items():
        for row in range(3, ws.max_row + 1):
            label = ws.cell(row, LABEL_COL_REF).value
            if not label:
                continue
            label_s = str(label).strip()
            val = parse_number(ws.cell(row, col).value)
            if val is not None:
                out[(label_s, col_key)] = val
    return out


def values_match(
    actual: float | None,
    expected: float | None,
    *,
    is_ratio: bool = False,
) -> bool:
    if actual is None and expected is None:
        return True
    if actual is None or expected is None:
        return False
    if is_ratio:
        tol = max(1e-6, abs(expected) * 0.002)
        return abs(actual - expected) <= tol
    tol = max(1.0, abs(expected) * 0.001)
    return abs(actual - expected) <= tol


def apply_template_validation_to_rows(
    rows: list[dict[str, Any]],
    sheet: str,
    year_list: list[int],
    *,
    is_ratio: bool = False,
    use_row_key: bool = False,
) -> dict[str, int]:
    """Disabled — extraction status comes from the report only, not xlsx comparison."""
    return {"matched": 0, "mismatched": 0, "compared": 0}


def apply_quarterly_template_validation(
    rows: list[dict[str, Any]],
    column_keys: list[str],
) -> dict[str, int]:
    """Disabled — extraction status comes from the report only, not xlsx comparison."""
    return {"matched": 0, "mismatched": 0, "compared": 0}


def _apply_template_validation_to_rows_legacy(
    rows: list[dict[str, Any]],
    sheet: str,
    year_list: list[int],
    *,
    is_ratio: bool = False,
    use_row_key: bool = False,
) -> dict[str, int]:
    """
    Compare preview row values to COMB model xlsx reference.
    Sets status template_mismatch and reference_values where they differ.
    Returns counts {matched, mismatched, compared}.
    """
    ref_by_row = reference_by_row(sheet, years=year_list)
    ref_by_label = reference_by_label(sheet, years=year_list)
    matched = mismatched = compared = 0

    for row in rows:
        if row.get("kind") != "data":
            continue
        label = str(row.get("label") or "")
        template_row = row.get("row") or row.get("drivers_row")
        values = row.setdefault("values", {})
        statuses = row.setdefault("statuses", {})
        references = row.setdefault("reference_values", {})

        for year in year_list:
            if year not in PILOT_YEARS:
                continue
            yk = str(year)
            actual = values.get(yk)
            expected: float | None = None
            if use_row_key and template_row:
                expected = ref_by_row.get((int(template_row), year))
            if expected is None:
                expected = ref_by_label.get((label, year))

            if expected is None:
                continue
            references[yk] = expected
            compared += 1
            if values_match(actual, expected, is_ratio=is_ratio):
                matched += 1
                if actual is not None and statuses.get(yk) not in (
                    STATUS_TEMPLATE_MISMATCH,
                ):
                    statuses[yk] = STATUS_FILLED
            else:
                mismatched += 1
                statuses[yk] = STATUS_TEMPLATE_MISMATCH

    return {
        "matched": matched,
        "mismatched": mismatched,
        "compared": compared,
    }


def attach_template_rows(
    preview_rows: list[dict[str, Any]],
    sheet: str,
) -> None:
    """Add template row numbers to preview rows for Drivers-style validation."""
    wb = _load_workbook_data_only()
    if sheet not in wb.sheetnames:
        return
    ws = wb[sheet]
    template_rows = extract_template_rows(ws)
    by_label: dict[str, int] = {}
    for item in template_rows:
        if item.get("kind") == "data":
            by_label[str(item["label"])] = int(item["row"])
    for row in preview_rows:
        if row.get("kind") != "data":
            continue
        lbl = str(row.get("label") or "")
        if lbl in by_label:
            row["row"] = by_label[lbl]
