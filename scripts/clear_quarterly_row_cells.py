"""
Clear quarterly P&L cell values in comb_workbook_data (values only — row stays).

By default clears the last row on the Quarterly sheet for year 2025 (all quarters).

Examples:
  python clear_quarterly_row_cells.py
  python clear_quarterly_row_cells.py --year 2025 --last-row
  python clear_quarterly_row_cells.py --label "Earnings per share" --year 2025
  python clear_quarterly_row_cells.py --dry-run
"""
from __future__ import annotations

import argparse
from datetime import datetime, timezone

import openpyxl

from comb_cell_status import STATUS_PENDING
from comb_workbook_store import COMB_COLLECTION, get_db, resolve_template, upsert_cells
from generate_comb_model import COMMERCIAL_BANK_SLUG, QUARTERLY_PILOT_QUARTERS


def quarterly_template_rows() -> list[dict[str, str | int]]:
    path = resolve_template()
    wb = openpyxl.load_workbook(path, data_only=False)
    ws = wb["Quarterly"]
    rows: list[dict[str, str | int]] = []
    for r in range(3, ws.max_row + 1):
        label = ws.cell(r, 2).value
        if label and str(label).strip():
            rows.append({"label": str(label).strip(), "row": r})
    wb.close()
    return rows


def clear_quarterly_cells(
    *,
    company_slug: str,
    label: str,
    year: int,
    quarters: list[str] | None = None,
    dry_run: bool = False,
) -> dict:
    quarter_list = quarters or list(QUARTERLY_PILOT_QUARTERS)
    db, client = get_db()
    coll = db[COMB_COLLECTION]
    docs: list[dict] = []
    cleared: list[dict] = []
    missing: list[str] = []

    for quarter in quarter_list:
        filt = {
            "company_slug": company_slug,
            "year": year,
            "report_type": "quarterly",
            "quarter": quarter,
            "sheet": "Quarterly",
            "label": label,
        }
        existing = coll.find_one(filt)
        if not existing:
            missing.append(quarter)
            continue
        old = existing.get("value")
        if old is None and existing.get("status") == STATUS_PENDING:
            continue
        cleared.append({"quarter": quarter, "old_value": old})
        if dry_run:
            continue
        doc = dict(existing)
        doc["value"] = None
        doc["status"] = STATUS_PENDING
        doc["updated_at"] = datetime.now(timezone.utc)
        docs.append(doc)

    updated = 0
    if docs and not dry_run:
        updated = upsert_cells(db, docs)

    client.close()
    return {
        "company_slug": company_slug,
        "label": label,
        "year": year,
        "cleared": cleared,
        "missing": missing,
        "updated": updated,
        "dry_run": dry_run,
    }


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--company-slug", default=COMMERCIAL_BANK_SLUG)
    ap.add_argument(
        "--label",
        default="",
        help="Row label to clear (default: last row on Quarterly sheet)",
    )
    ap.add_argument(
        "--last-row",
        action="store_true",
        help="Clear the last row on the Quarterly sheet (default when --label omitted)",
    )
    ap.add_argument("--year", type=int, default=2025, help="Calendar year (default: 2025)")
    ap.add_argument(
        "--quarters",
        default="",
        help="Comma-separated quarters to clear, e.g. Q1,Q2 (default: all Q1–Q4)",
    )
    ap.add_argument(
        "--dry-run",
        action="store_true",
        help="Show what would be cleared without writing to MongoDB",
    )
    args = ap.parse_args()

    template_rows = quarterly_template_rows()
    if not template_rows:
        print("No rows found on Quarterly sheet.")
        return 1

    label = args.label.strip()
    if not label:
        label = str(template_rows[-1]["label"])
        print(f"Using last Quarterly row: {label!r}")

    quarters = (
        [q.strip().upper() for q in args.quarters.split(",") if q.strip()]
        if args.quarters.strip()
        else None
    )

    result = clear_quarterly_cells(
        company_slug=args.company_slug.strip(),
        label=label,
        year=args.year,
        quarters=quarters,
        dry_run=args.dry_run,
    )

    print(f"Company: {result['company_slug']}")
    print(f"Label:   {result['label']}")
    print(f"Year:    {result['year']}")
    if result["dry_run"]:
        print("Mode:    dry-run (no changes written)")
    for item in result["cleared"]:
        print(f"  {item['quarter']}: clear {item['old_value']!r}")
    if not result["cleared"]:
        print("  (no filled cells to clear)")
    for q in result["missing"]:
        print(f"  {q}: no document in DB")
    if not result["dry_run"]:
        print(f"Updated {result['updated']} cell(s) in {COMB_COLLECTION}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
